from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / ".src" / "maven-project"
OUT = ROOT / "TranscriptFetch Java Study Guide.docx"

NAVY = "18324A"
BLUE = "2563A6"
TEAL = "1C7C7D"
PALE = "EAF2F8"
LIGHT = "F4F6F8"
MID = "D7E0E8"
DARK = RGBColor(31, 45, 61)
MUTED = RGBColor(83, 101, 116)


def shade(cell_or_p, fill):
    element = cell_or_p._tc if hasattr(cell_or_p, "_tc") else cell_or_p._p
    props_tag = "w:tcPr" if hasattr(cell_or_p, "_tc") else "w:pPr"
    props = element.get_or_add_tcPr() if hasattr(cell_or_p, "_tc") else element.get_or_add_pPr()
    shd = props.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        props.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_rule(doc, color=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def add_callout(doc, title, text, fill=PALE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.75)
    shade(cell, fill)
    set_cell_margins(cell, 140, 180, 140, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor(24, 50, 74)
    p2 = cell.add_paragraph(text)
    p2.paragraph_format.space_after = Pt(0)
    return table


def add_code(doc, text, numbered=False, font_size=7.4):
    lines = text.rstrip("\n").splitlines() or [""]
    width = len(str(len(lines)))
    for i, line in enumerate(lines, 1):
        shown = f"{i:>{width}}  {line}" if numbered else line
        p = doc.add_paragraph(style="Code")
        p.paragraph_format.keep_together = False
        p.paragraph_format.keep_with_next = False
        shade(p, LIGHT)
        run = p.add_run(shown)
        run.font.size = Pt(font_size)


def add_bullet(doc, lead, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(lead).bold = True
    p.add_run(text)
    return p


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    return h


def add_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    hdr = t.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        c = hdr.cells[i]
        c.text = text
        shade(c, NAVY)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs:
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.bold = True
            r.font.size = Pt(9)
        set_cell_margins(c)
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = str(text)
            if len(t.rows) % 2 == 1:
                shade(cells[i], "F7F9FB")
            set_cell_margins(cells[i])
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(8.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    return t


def configure(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.62)
    sec.bottom_margin = Inches(0.62)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)
    sec.header_distance = Inches(0.28)
    sec.footer_distance = Inches(0.28)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(34)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor(24, 50, 74)

    for name, size, color in (("Heading 1", 20, NAVY), ("Heading 2", 14, BLUE), ("Heading 3", 11, TEAL)):
        s = styles[name]
        s.font.name = "Aptos Display"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(12)
        s.paragraph_format.space_after = Pt(5)

    if "Code" not in [s.name for s in styles]:
        code = styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code"]
    code.font.name = "DejaVu Sans Mono"
    code.font.size = Pt(7.4)
    code.font.color.rgb = RGBColor(31, 41, 55)
    code.paragraph_format.left_indent = Inches(0.10)
    code.paragraph_format.right_indent = Inches(0.05)
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 0.88

    header = sec.header.paragraphs[0]
    header.text = "TRANSCRIPTFETCH  /  JAVA STUDY GUIDE"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in header.runs:
        r.font.name = "Aptos"
        r.font.size = Pt(8)
        r.font.bold = True
        r.font.color.rgb = MUTED
    add_page_number(sec.footer.paragraphs[0])


def build():
    production = (SRC / "src/main/java/TranscriptFetch.java").read_text()
    pom = (SRC / "pom.xml").read_text()
    tests = (SRC / "src/test/java/TranscriptFetchTest.java").read_text()

    doc = Document()
    configure(doc)
    doc.core_properties.title = "TranscriptFetch Java Study Guide"
    doc.core_properties.subject = "Production Java, Maven, and JUnit walkthrough"
    doc.core_properties.author = "Muse"

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(58)
    r = p.add_run("TRANSCRIPTFETCH")
    r.font.name = "Aptos Display"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(TEAL)
    title = doc.add_paragraph(style="Title")
    title.add_run("Java Study Guide")
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(18)
    rs = sub.add_run("A real production transcript fetcher, explained from command line to unit test")
    rs.font.size = Pt(15)
    rs.font.color.rgb = MUTED
    add_rule(doc)
    add_callout(doc, "What you get", "A guided code tour, a standard Maven project, seven deterministic JUnit 5 tests, practice exercises, and the complete source listings.")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(150)
    r = p.add_run("Written by Muse  |  September 2026")
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED
    doc.add_page_break()

    add_heading(doc, "Start with the shape of the program", 1)
    doc.add_paragraph("TranscriptFetch accepts one or more YouTube video IDs. For each ID, it first asks the shared cache. A cache miss triggers a player-metadata request, English-caption selection, caption download, XML cleanup, and a cache write. The command prints one small JSON result per video so another job can consume it safely.")
    add_table(doc, ["Stage", "Input", "Output"], [
        ("1. Cache check", "video ID", "record or miss"),
        ("2. Player request", "video ID", "metadata + tracks"),
        ("3. Track choice", "caption tracks", "manual or auto English"),
        ("4. Caption parse", "timed-text XML", "plain transcript"),
        ("5. Cache write", "full record", "durable JSON"),
        ("6. Result", "status fields", "one JSON line"),
    ], [1.45, 2.15, 2.95])
    add_callout(doc, "Design idea", "The cache sits before the network. That is both a speed optimization and a reliability feature: a transcript already fetched does not depend on YouTube being reachable again.", "E8F5F3")

    add_heading(doc, "Project map", 2)
    add_code(doc, "transcript-fetch/\n├── pom.xml\n└── src/\n    ├── main/java/TranscriptFetch.java\n    └── test/java/TranscriptFetchTest.java", font_size=8.2)
    doc.add_paragraph("Maven recognizes this directory layout automatically. Production classes live under src/main/java; test classes live under src/test/java. The pom.xml file supplies the compiler settings, dependencies, and test runner.")

    add_heading(doc, "Read main as the command-line contract", 1)
    doc.add_paragraph("The main method validates the first token, collects every video ID, accepts an optional job label, then processes IDs independently. A failed video does not prevent later IDs from being attempted. The final exit code summarizes the batch: 0 means all resolved; 2 means at least one failed.")
    add_bullet(doc, "args[0]", " must equal fetch. This gives the script room to add other commands later.")
    add_bullet(doc, "List<String> ids", " holds an arbitrary batch rather than forcing one process per video.")
    add_bullet(doc, "try/catch per ID", " converts exceptions into structured JSON instead of ending the entire batch.")
    add_bullet(doc, "anyFail", " preserves the batch result after every ID has been attempted.")
    add_callout(doc, "Study question", "Why print one JSON object per line instead of one large JSON array? Hint: think about streaming output and partial progress.")

    add_heading(doc, "fetchOne is the orchestration method", 1)
    doc.add_paragraph("The method fetchOne reads like a checklist. It delegates each technical concern to a smaller method, which keeps the high-level flow visible.")
    add_table(doc, ["Method", "Responsibility", "Side effect"], [
        ("cacheGet", "lookup + freshness", "hit log"),
        ("postPlayer", "metadata + tracks", "HTTP POST"),
        ("curlGet", "caption XML", "HTTP GET"),
        ("parseCaptions", "XML-ish text cleanup", "none"),
        ("cachePut", "persist record", "file write via helper"),
    ], [1.35, 2.45, 2.75])
    doc.add_paragraph("Notice the order of manual and automatic captions. The loop records the first manual English track and the first automatic-speech-recognition track. The later conditional chooses manual when available, otherwise automatic. This is a compact policy encoded in ordinary variables rather than a complicated comparator.")

    add_heading(doc, "JSON traversal stays defensive", 2)
    doc.add_paragraph("Responses from external services can omit nested objects. The code uses optJSONObject, optJSONArray, and optString for optional values, then uses getString only after a caption track has been selected. This separates uncertain input from the field the next operation requires.")
    add_bullet(doc, "opt... methods", " return a default or null instead of throwing when a field is absent.")
    add_bullet(doc, "getString(\"baseUrl\")", " is intentionally strict because a selected caption track must have a URL.")
    add_bullet(doc, "JSONObject.NULL", " preserves a true JSON null rather than writing the Java string \"null\".")

    add_heading(doc, "The network layer is isolated behind one function", 1)
    doc.add_paragraph("All external requests eventually reach curl. ProcessBuilder constructs the argument list without invoking a shell, captures standard output and standard error separately, waits for completion, then interprets the final line as an HTTP status. Returning the response body only for HTTP 200 gives callers a simple contract.")
    add_callout(doc, "Why ProcessBuilder matters", "A List<String> preserves argument boundaries. The URL remains one argument, and the code avoids shell quoting rules. This is safer and easier to reason about than building one command string.")
    add_bullet(doc, "-sS", " keeps normal output quiet while preserving errors.")
    add_bullet(doc, "-m 60", " bounds a request so a batch cannot hang forever on one network call.")
    add_bullet(doc, "-w", " appends the HTTP status after the body, making status handling explicit.")

    add_heading(doc, "Retry only failures that may recover", 2)
    doc.add_paragraph("withBackoff accepts behavior through the CurlCall functional interface. It retries HTTP 429, server errors from 500 upward, and other transient exceptions. A non-retryable HTTP error such as 404 is thrown immediately. The delay doubles with each attempt and adds random jitter so repeated jobs are less likely to collide again.")
    add_table(doc, ["Failure", "Retry?", "Reason"], [
        ("HTTP 429", "Yes", "temporary rate limit"),
        ("HTTP 5xx", "Yes", "temporary server fault"),
        ("Network error", "Yes", "may be transient"),
        ("HTTP 404", "No", "same request unlikely to help"),
    ], [1.6, 1.0, 3.95])
    doc.add_paragraph("The test suite deliberately checks the 404 rule because retry policies can accidentally turn a fast permanent failure into a slow repeated failure.")

    add_heading(doc, "Caption parsing supports two YouTube formats", 1)
    doc.add_paragraph("parseCaptions first looks for <s> elements used by one timed-text format. If none are found, it falls back to <text> elements used by another format. Each cue passes through appendCue, which removes nested tags, decodes the supported entities, collapses whitespace, and joins non-empty cues with one space.")
    add_table(doc, ["Input detail", "Transformation", "Example"], [
        ("nested tag", "remove markup", "<font>now</font> -> now"),
        ("entity", "decode text", "&amp; -> &"),
        ("line breaks", "collapse whitespace", "many\\n spaces -> many spaces"),
        ("blank cue", "skip", "<text>   </text> -> nothing"),
    ], [1.6, 2.0, 2.95])
    add_callout(doc, "Important boundary", "This parser handles the caption formats the fetcher receives; it is not a general-purpose XML parser. If the format becomes more complex, a real XML parser would be the safer next step.", "FFF4E5")

    add_heading(doc, "Maven turns the folder into a repeatable build", 1)
    doc.add_paragraph("The Project Object Model, pom.xml, is Maven's build description. The coordinates identify the project, properties establish Java 21 and UTF-8, dependencies bring in JSON support and JUnit, and the Surefire plugin discovers and runs test classes during the test phase.")
    add_table(doc, ["POM element", "Value", "Purpose"], [
        ("groupId", "com.muse.youtube", "project namespace"),
        ("artifactId", "transcript-fetch", "project name"),
        ("compiler.release", "21", "Java language/API level"),
        ("org.json", "20240303", "JSON objects + arrays"),
        ("JUnit Jupiter", "5.11.4", "test API + engine"),
        ("Surefire", "3.5.2", "runs tests in Maven"),
    ], [1.55, 1.8, 3.2])
    add_heading(doc, "Commands to practice", 2)
    add_code(doc, "# Tell Maven which JDK to use\nexport JAVA_HOME=\"$HOME/workspace/java/jdk-21\"\n\n# Compile production + tests, then run all tests\nmvn test\n\n# Remove old build output and test from a clean state\nmvn clean test\n\n# Build target/transcript-fetch-1.0.0.jar after tests pass\nmvn package", font_size=8.0)
    doc.add_paragraph("Maven phases are ordered. package includes compile and test; clean deletes the target directory before a new lifecycle begins. A dependency marked test is available to test code but is not treated as a production dependency.")

    add_heading(doc, "The tests target deterministic seams", 1)
    doc.add_paragraph("The seven JUnit tests avoid live YouTube calls. Network tests would be slower, less repeatable, and vulnerable to rate limits. Instead, the suite verifies the pure parser and the retry classification rule. This is the right first layer: fast tests that explain exactly which behavior broke.")
    add_table(doc, ["Test", "Behavior proved", "Network"], [
        ("parsesFormat3Cues", "joins <s> cues", "none"),
        ("parsesSrv1TextCues", "fallback to <text>", "none"),
        ("decodesEntities...", "text cleanup", "none"),
        ("normalizesWhitespace...", "spaces + blanks", "none"),
        ("returnsEmptyString...", "empty input", "none"),
        ("withBackoffReturns...", "first-call success", "none"),
        ("withBackoffDoesNot...", "404 fails once", "none"),
    ], [2.15, 3.15, 1.25])
    add_callout(doc, "Verified result", "The companion JUnit 5 suite completed with 7 tests successful and 0 failed. The tests were executed directly with the JUnit console because this runtime blocks Maven's Java-based dependency downloads; the same test class is wired for mvn test on a normal Maven-enabled machine.", "E8F5F3")

    add_heading(doc, "Use Arrange, Act, Assert when reading a test", 2)
    doc.add_paragraph("Each test is intentionally split into three visual blocks. Arrange creates the input. Act calls one production method. Assert compares the actual result with the expected result. Keeping one behavior per test makes a failure message useful.")
    add_code(doc, "String xml = \"<transcript><text>Hello</text></transcript>\"; // Arrange\nString result = TranscriptFetch.parseCaptions(xml);              // Act\nassertEquals(\"Hello\", result);                                 // Assert", font_size=8.0)

    doc.add_page_break()
    add_heading(doc, "Practice in small, safe steps", 1)
    add_bullet(doc, "Exercise 1 - add &apos; decoding.", " Write a failing test first, then add one replacement in appendCue. Run mvn test.")
    add_bullet(doc, "Exercise 2 - test HTTP 500 retry.", " The current method sleeps, so first refactor delay calculation behind an injectable sleeper. Your test should avoid real waiting.")
    add_bullet(doc, "Exercise 3 - prefer a named language.", " Add a language argument and test track selection separately from the network response.")
    add_bullet(doc, "Exercise 4 - replace regex parsing.", " Use an XML parser and rerun the existing tests unchanged. Green tests show preserved behavior.")
    add_bullet(doc, "Exercise 5 - add integration boundaries.", " Stub the player response and cache process so fetchOne can be tested without external systems.")
    add_callout(doc, "Best learning loop", "Read one method, predict its output for one input, write a test for that prediction, run the test, then change only one thing.")

    add_heading(doc, "Full source: TranscriptFetch.java", 1)
    doc.add_paragraph("This is the current production source used by the YouTube transcript workflow. Line numbers are added only for reading; they are not part of the file.")
    add_code(doc, production, numbered=True, font_size=6.9)

    add_heading(doc, "Full source: pom.xml", 1)
    doc.add_paragraph("This build file wraps the production class in a standard Java 21 Maven project.")
    add_code(doc, pom, numbered=True, font_size=6.8)

    add_heading(doc, "Full source: TranscriptFetchTest.java", 1)
    doc.add_paragraph("These are the seven deterministic JUnit 5 tests described earlier.")
    add_code(doc, tests, numbered=True, font_size=6.7)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
